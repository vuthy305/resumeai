"""DOCX text extraction (application convenience; the LayoutLMv3 path uses PDF/images)."""
from pathlib import Path


class DOCXError(ValueError):
    pass


def extract_docx_text(path: str | Path) -> str:
    try:
        import docx
        document = docx.Document(str(path))
    except Exception as e:
        raise DOCXError(f"Corrupted or unreadable DOCX: {e}") from e
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    text = "\n".join(l for l in lines if l.strip())
    if not text.strip():
        raise DOCXError("DOCX contains no text.")
    return text
