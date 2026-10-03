from datetime import date

import pytest
from PIL import Image

from src.document_processing.docx_processor import DOCXError, extract_docx_text
from src.document_processing.image_processor import ImagePreprocessor
from src.document_processing.pdf_processor import PDFError, extract_text_layer_words, iter_pdf_pages
from src.config import PreprocessConfig
from src.extraction.baseline import RuleBasedExtractor, is_section_header
from src.extraction.entities import bio_to_entities, entities_to_fields
from src.extraction.metrics import seqeval_rows
from src.extraction.profile_builder import ProfileBuilder, extract_experience
from src.pipeline import CVPipeline, PipelineError
from src.utils.text_utils import find_emails, find_phones


def test_email_extraction():
    assert find_emails("Contact: jane.doe+cv@mail.co.uk.") == ["jane.doe+cv@mail.co.uk"]
    assert find_emails("no email here") == []


def test_phone_extraction():
    assert find_phones("Tel: +855 12 345 678") == ["+855 12 345 678"]
    assert find_phones("(555) 123-4567") == ["(555) 123-4567"]
    assert find_phones("2018 - 2021") == []          # year range is not a phone
    assert find_phones("12/05/2020") == []


def _page():
    words = ["Jane", "Doe", "12", "Main", "Street,", "Springfield", "jane@doe.com", "+1", "555", "123", "4567",
             "EDUCATION", "BSc", "Physics", "WORK", "EXPERIENCE", "Analyst"]
    y = [50, 50, 90, 90, 90, 90, 120, 120, 120, 120, 120, 400, 430, 430, 500, 500, 530]
    x = [100, 200, 100, 140, 200, 300, 100, 400, 430, 470, 510, 100, 100, 160, 100, 180, 100]
    boxes = [[a, b, a + 50, b + 15] for a, b in zip(x, y)]
    return words, boxes


def test_baseline_tags_all_entity_types():
    words, boxes = _page()
    labels = RuleBasedExtractor().tag(words, boxes)
    f = entities_to_fields(bio_to_entities(words, labels))
    assert f["name"] == "Jane Doe"
    assert f["email"] == "jane@doe.com"
    assert f["phone"] == "+1 555 123 4567"
    assert "Main Street" in f["address"]
    assert f["section_headers"] == ["EDUCATION", "WORK EXPERIENCE"]


def test_section_header_rules():
    assert is_section_header("Work Experience:")
    assert not is_section_header("Skills: Python, SQL, Docker")


def test_entity_reconstruction_handles_orphan_I():
    ents = bio_to_entities(["A", "B", "C", "D"], ["B-NAME", "I-NAME", "O", "I-EMAIL"], page=2)
    assert [(e["entity"], e["text"], e["page"]) for e in ents] == [("NAME", "A B", 2), ("EMAIL", "D", 2)]


def test_seqeval_metrics_perfect_and_macro():
    rows = seqeval_rows([["B-NAME", "I-NAME", "O", "B-EMAIL"]], [["B-NAME", "I-NAME", "O", "O"]])
    d = {r["entity"]: r for r in rows}
    assert d["NAME"]["f1"] == 100.0 and d["EMAIL"]["f1"] == 0.0 and d["MACRO"]["f1"] == 50.0


def test_profile_parsing():
    text = """SKILLS
Python, Postgres, Docker
WORK EXPERIENCE
Backend Developer
Acme Technology Co., Ltd | Jan 2020 - Dec 2022
EDUCATION
Master of Science in Computer Science
LANGUAGES
English, Khmer"""
    p = ProfileBuilder().build(text, {"name": "X"}, today=date(2026, 1, 1))
    assert p["skills"] == ["Python", "PostgreSQL", "Docker"]
    assert p["highest_degree"] == "Master's Degree" and p["highest_degree_level"] == 4
    assert p["total_experience_years"] == 2.9          # Jan 2020 -> Dec 2022 = 35 months
    assert p["job_titles"] == ["Backend Developer"]
    assert "Acme Technology Co., Ltd" in p["companies"]
    assert p["languages"] == ["English", "Khmer"]


def test_experience_overlaps_merged_and_explicit_statement():
    _, years, src = extract_experience(["2019 - 2021", "2020 - 2022"], today=date(2026, 1, 1))
    assert years == 3.0 and src == "date_ranges"
    _, years, src = extract_experience(["5+ years of experience in banking"])
    assert years == 5.0 and src == "explicit_statement"
    _, years, src = extract_experience(["No dates here"])
    assert years == 0 and src == "not_found"


def test_pdf_processing(sample_pdf, tmp_path):
    pages = list(iter_pdf_pages(sample_pdf, dpi=72))
    assert len(pages) == 1 and pages[0][0] == 1 and pages[0][1].mode == "RGB"
    layer = extract_text_layer_words(sample_pdf)[0]
    assert "Dara" in layer["words"] and all(0 <= v <= 1000 for b in layer["boxes"] for v in b)
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"%PDF-1.4 garbage")
    with pytest.raises(PDFError):
        list(iter_pdf_pages(bad))


def test_docx_processing(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Jane Doe")
    d.add_paragraph("jane@doe.com")
    d.save(tmp_path / "cv.docx")
    assert "jane@doe.com" in extract_docx_text(tmp_path / "cv.docx")
    empty = docx.Document()
    empty.save(tmp_path / "empty.docx")
    with pytest.raises(DOCXError):
        extract_docx_text(tmp_path / "empty.docx")


def test_image_preprocessing_is_configurable():
    img = Image.new("RGB", (4000, 1000), "white")
    assert max(ImagePreprocessor()(img).size) == 2000
    out = ImagePreprocessor(PreprocessConfig(grayscale=True, contrast=True, denoise=True, sharpen=True))(img)
    assert out.mode == "RGB"


def test_pipeline_pdf_fallback(sample_pdf):
    r = CVPipeline().process_file(sample_pdf)
    assert r["extraction_method"] == "rule_based_baseline" and r["ocr_engine"] == "pdf_text_layer"
    p = r["profile"]
    assert p["name"] == "Dara Vann" and p["email"] == "dara.vann@example.com"
    assert {"Python", "PostgreSQL", "REST API", "Docker", "FastAPI"} <= set(p["skills"])
    assert p["highest_degree_level"] == 3


def test_pipeline_rejects_unsupported_and_images_without_ocr(tmp_path):
    f = tmp_path / "cv.txt"
    f.write_text("hello")
    with pytest.raises(PipelineError):
        CVPipeline().process_file(f)
    img = tmp_path / "cv.png"
    Image.new("RGB", (100, 100), "white").save(img)
    with pytest.raises(PipelineError):
        CVPipeline().process_file(img)
